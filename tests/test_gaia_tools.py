from pathlib import Path

from eval.gaia.harness import failure_mode, final_answer
from eval.gaia.scoring import exact_match, summarize
from eval.gaia.tools import calculate, read_file


def test_calculator_executes_restricted_arithmetic():
    assert calculate("(7 + 5) * 3").text == "36"
    assert calculate("__import__('os').system('echo unsafe')").error == "calculator_error"


def test_file_reader_executes_and_confines_paths(tmp_path: Path):
    attachment = tmp_path / "evidence.csv"
    attachment.write_text("name,value\nalpha,42\n")
    result = read_file("evidence.csv", tmp_path)
    assert result.error is None
    assert "alpha\t42" in result.text
    assert read_file("../secret.txt", tmp_path).error == "file_access_error"


def test_attachment_readers_execute_real_parsers(tmp_path: Path):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    from pypdf import PdfWriter

    workbook = Workbook()
    workbook.active["A1"] = "spreadsheet evidence"
    workbook.save(tmp_path / "evidence.xlsx")
    assert "spreadsheet evidence" in read_file("evidence.xlsx", tmp_path).text

    document = Document()
    document.add_paragraph("document evidence")
    document.save(tmp_path / "evidence.docx")
    assert "document evidence" in read_file("evidence.docx", tmp_path).text

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with (tmp_path / "evidence.pdf").open("wb") as handle:
        writer.write(handle)
    assert read_file("evidence.pdf", tmp_path).error is None

    Image.new("RGB", (2, 2), "white").save(tmp_path / "evidence.png")
    image = read_file("evidence.png", tmp_path)
    assert image.error is None
    assert image.image_media_type == "image/png"
    assert image.image_base64


def test_exact_match_and_wilson_summary():
    assert exact_match(" 1,234 ", "1234")
    result = summarize(
        [
            {"correct": True, "failure_mode": "none"},
            {"correct": False, "failure_mode": "reasoning_error"},
        ]
    )
    assert result["accuracy"]["events"] == 1
    assert result["accuracy"]["n_trials"] == 2
    assert result["accuracy"]["ci_95"][0] < 0.5 < result["accuracy"]["ci_95"][1]


def test_failure_attribution_and_final_answer():
    assert final_answer("work\nFINAL_ANSWER: Mars") == "Mars"
    assert (
        failure_mode(correct=False, attachment=True, trace=[])
        == "file_read_error"
    )
    assert (
        failure_mode(
            correct=False,
            attachment=False,
            trace=[{"tool": "web_search", "error": "retrieval_error"}],
        )
        == "retrieval_error"
    )
